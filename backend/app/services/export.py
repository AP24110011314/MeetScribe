import io
import json
from datetime import date, datetime, timezone
from typing import List, Optional

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.db.models import ActionItem, Meeting, Summary, Transcript


def format_duration(seconds: Optional[float]) -> str:
    if not seconds or seconds <= 0:
        return "Unknown"
    mins, secs = divmod(int(seconds), 60)
    hours, mins = divmod(mins, 60)
    if hours > 0:
        return f"{hours}h {mins}m {secs}s"
    return f"{mins}m {secs}s"


def export_markdown(
    meeting: Meeting,
    summary: Optional[Summary],
    action_items: List[ActionItem],
    transcript: Optional[Transcript] = None,
) -> str:
    """
    Renders meeting notes, summary, decisions, and action items as Markdown.
    """
    date_str = (
        meeting.created_at.strftime("%Y-%m-%d %H:%M UTC")
        if meeting.created_at
        else datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    )
    duration_str = format_duration(meeting.duration_sec)

    lines = [
        f"# {meeting.title}",
        "",
        f"- **Date:** {date_str}",
        f"- **Duration:** {duration_str}",
        f"- **Language:** {meeting.language or 'en'}",
        "",
    ]

    if summary:
        lines.extend(
            [
                "## Executive Summary",
                "",
                summary.summary,
                "",
            ]
        )

        if summary.key_points:
            lines.extend(
                [
                    "## Key Discussion Points",
                    "",
                ]
            )
            for point in summary.key_points:
                lines.append(f"- {point}")
            lines.append("")

        if summary.decisions:
            lines.extend(
                [
                    "## Decisions",
                    "",
                ]
            )
            for decision in summary.decisions:
                lines.append(f"- {decision}")
            lines.append("")

    if action_items:
        lines.extend(
            [
                "## Action Items",
                "",
                "| Description | Owner | Due Date | Priority | Status |",
                "|---|---|---|---|---|",
            ]
        )
        for item in action_items:
            owner_str = item.owner or "Unassigned"
            due_str = item.due_date.isoformat() if item.due_date else "None"
            priority_str = item.priority.capitalize()
            status_str = item.status.replace("_", " ").capitalize()
            desc = item.description.replace("|", "/")
            lines.append(f"| {desc} | {owner_str} | {due_str} | {priority_str} | {status_str} |")
        lines.append("")

    if transcript and transcript.full_text:
        lines.extend(
            [
                "## Transcript",
                "",
                transcript.full_text,
                "",
            ]
        )

    return "\n".join(lines)


def export_pdf(
    meeting: Meeting,
    summary: Optional[Summary],
    action_items: List[ActionItem],
) -> bytes:
    """
    Renders meeting notes as a styled PDF using ReportLab.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=10,
    )
    heading2_style = ParagraphStyle(
        "DocH2",
        parent=styles["Heading2"],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#334155"),
        spaceBefore=14,
        spaceAfter=6,
    )
    meta_style = ParagraphStyle(
        "DocMeta",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=14,
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=6,
    )
    bullet_style = ParagraphStyle(
        "DocBullet",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1e293b"),
        leftIndent=15,
        spaceAfter=4,
    )

    story = []

    # Title & Metadata
    story.append(Paragraph(meeting.title, title_style))
    date_str = (
        meeting.created_at.strftime("%B %d, %Y at %H:%M UTC")
        if meeting.created_at
        else datetime.now(timezone.utc).strftime("%B %d, %Y")
    )
    meta_text = f"Date: {date_str}  •  Duration: {format_duration(meeting.duration_sec)}  •  Language: {meeting.language or 'en'}"
    story.append(Paragraph(meta_text, meta_style))
    story.append(Spacer(1, 10))

    # Executive Summary
    if summary:
        story.append(Paragraph("Executive Summary", heading2_style))
        story.append(Paragraph(summary.summary, body_style))
        story.append(Spacer(1, 8))

        # Key Points
        if summary.key_points:
            story.append(Paragraph("Key Points", heading2_style))
            for kp in summary.key_points:
                story.append(Paragraph(f"• {kp}", bullet_style))
            story.append(Spacer(1, 8))

        # Decisions
        if summary.decisions:
            story.append(Paragraph("Decisions", heading2_style))
            for d in summary.decisions:
                story.append(Paragraph(f"• {d}", bullet_style))
            story.append(Spacer(1, 8))

    # Action Items Table
    if action_items:
        story.append(Paragraph("Action Items", heading2_style))
        table_data = [["Description", "Owner", "Due Date", "Priority", "Status"]]
        for item in action_items:
            owner_text = item.owner or "—"
            due_text = item.due_date.isoformat() if item.due_date else "—"
            table_data.append(
                [
                    Paragraph(item.description, body_style),
                    owner_text,
                    due_text,
                    item.priority.capitalize(),
                    item.status.replace("_", " ").capitalize(),
                ]
            )

        col_widths = [200, 75, 75, 60, 65]
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 9),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(t)

    doc.build(story)
    return buffer.getvalue()


def export_docx(
    meeting: Meeting,
    summary: Optional[Summary],
    action_items: List[ActionItem],
    transcript: Optional[Transcript] = None,
) -> bytes:
    """
    Renders meeting notes, summary, and action items as a styled Microsoft Word DOCX document.
    """
    doc = docx.Document()

    # Document Title
    doc.add_heading(meeting.title, level=1)

    # Metadata
    date_str = (
        meeting.created_at.strftime("%B %d, %Y at %H:%M UTC")
        if meeting.created_at
        else "Unknown date"
    )
    meta_p = doc.add_paragraph()
    meta_p.add_run(
        f"Date: {date_str}  |  Duration: {format_duration(meeting.duration_sec)}  |  Language: {meeting.language or 'en'}"
    ).italic = True

    # Executive Summary
    if summary:
        doc.add_heading("Executive Summary", level=2)
        doc.add_paragraph(summary.summary)

        if summary.key_points:
            doc.add_heading("Key Points", level=2)
            for kp in summary.key_points:
                doc.add_paragraph(kp, style="List Bullet")

        if summary.decisions:
            doc.add_heading("Decisions", level=2)
            for d in summary.decisions:
                doc.add_paragraph(d, style="List Bullet")

    # Action Items Table
    if action_items:
        doc.add_heading("Action Items", level=2)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        hdr_titles = ["Description", "Owner", "Due Date", "Priority", "Status"]
        for i, h_text in enumerate(hdr_titles):
            hdr_cells[i].text = h_text
            shd = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls("w")))
            hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
            if hdr_cells[i].paragraphs[0].runs:
                hdr_cells[i].paragraphs[0].runs[0].bold = True

        for item in action_items:
            row_cells = table.add_row().cells
            row_cells[0].text = item.description
            row_cells[1].text = item.owner or "—"
            row_cells[2].text = item.due_date.isoformat() if item.due_date else "—"
            row_cells[3].text = item.priority.capitalize()
            row_cells[4].text = item.status.replace("_", " ").capitalize()

    # Optional transcript
    if transcript and transcript.full_text:
        doc.add_heading("Transcript", level=2)
        doc.add_paragraph(transcript.full_text)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_txt(
    meeting: Meeting,
    summary: Optional[Summary],
    action_items: List[ActionItem],
    transcript: Optional[Transcript] = None,
) -> str:
    """
    Renders meeting notes as clean plain text with headers and separators.
    """
    lines = [
        "=" * 60,
        meeting.title.upper(),
        "=" * 60,
    ]
    date_str = (
        meeting.created_at.strftime("%Y-%m-%d %H:%M UTC") if meeting.created_at else "Unknown date"
    )
    lines.append(f"Date: {date_str}")
    lines.append(f"Duration: {format_duration(meeting.duration_sec)}")
    lines.append(f"Language: {meeting.language or 'en'}")
    lines.append("")

    if summary:
        lines.append("-" * 40)
        lines.append("EXECUTIVE SUMMARY")
        lines.append("-" * 40)
        lines.append(summary.summary)
        lines.append("")

        if summary.key_points:
            lines.append("-" * 40)
            lines.append("KEY POINTS")
            lines.append("-" * 40)
            for kp in summary.key_points:
                lines.append(f"• {kp}")
            lines.append("")

        if summary.decisions:
            lines.append("-" * 40)
            lines.append("DECISIONS")
            lines.append("-" * 40)
            for d in summary.decisions:
                lines.append(f"• {d}")
            lines.append("")

    if action_items:
        lines.append("-" * 40)
        lines.append("ACTION ITEMS")
        lines.append("-" * 40)
        for item in action_items:
            owner = item.owner or "Unassigned"
            due = item.due_date.isoformat() if item.due_date else "None"
            lines.append(
                f"[{item.status.upper()}] {item.description} (Owner: {owner}, Due: {due}, Priority: {item.priority.capitalize()})"
            )
        lines.append("")

    if transcript and transcript.full_text:
        lines.append("-" * 40)
        lines.append("TRANSCRIPT")
        lines.append("-" * 40)
        lines.append(transcript.full_text)
        lines.append("")

    return "\n".join(lines)


def export_json(
    meeting: Meeting,
    summary: Optional[Summary],
    action_items: List[ActionItem],
    transcript: Optional[Transcript] = None,
) -> str:
    """
    Renders structured dump of meeting metadata, transcript, summary, and action items.
    """
    data = {
        "id": str(meeting.id),
        "title": meeting.title,
        "created_at": meeting.created_at.isoformat() if meeting.created_at else None,
        "duration_sec": meeting.duration_sec,
        "language": meeting.language,
        "summary": (
            {
                "summary": summary.summary,
                "key_points": summary.key_points,
                "decisions": summary.decisions,
            }
            if summary
            else None
        ),
        "action_items": [
            {
                "id": str(item.id),
                "description": item.description,
                "owner": item.owner,
                "due_date": item.due_date.isoformat() if item.due_date else None,
                "priority": item.priority,
                "status": item.status,
            }
            for item in action_items
        ],
        "transcript": (
            {
                "full_text": transcript.full_text,
                "segments": transcript.segments,
            }
            if transcript
            else None
        ),
    }
    return json.dumps(data, indent=2, ensure_ascii=False)


def _format_srt_time(seconds: float) -> str:
    ms = int(round((seconds - int(seconds)) * 1000))
    s = int(seconds) % 60
    m = (int(seconds) // 60) % 60
    h = int(seconds) // 3600
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def export_srt(transcript: Optional[Transcript]) -> str:
    """
    Renders SubRip (.srt) subtitle format from transcript segments.
    """
    if not transcript or not transcript.segments:
        return ""
    blocks = []
    for i, seg in enumerate(transcript.segments, start=1):
        start = seg.get("start", 0.0)
        end = seg.get("end", start + 2.0)
        text = seg.get("text", "").strip()
        time_line = f"{_format_srt_time(start)} --> {_format_srt_time(end)}"
        blocks.append(f"{i}\n{time_line}\n{text}")
    return "\n\n".join(blocks) + "\n"


def _format_vtt_time(seconds: float) -> str:
    ms = int(round((seconds - int(seconds)) * 1000))
    s = int(seconds) % 60
    m = (int(seconds) // 60) % 60
    h = int(seconds) // 3600
    return f"{h:02d}:{m:02d}:{s:02d}.{ms:03d}"


def export_vtt(transcript: Optional[Transcript]) -> str:
    """
    Renders WebVTT (.vtt) subtitle format from transcript segments.
    """
    lines = ["WEBVTT", ""]
    if not transcript or not transcript.segments:
        return "WEBVTT\n"
    for i, seg in enumerate(transcript.segments, start=1):
        start = seg.get("start", 0.0)
        end = seg.get("end", start + 2.0)
        text = seg.get("text", "").strip()
        time_line = f"{_format_vtt_time(start)} --> {_format_vtt_time(end)}"
        lines.append(f"{i}")
        lines.append(time_line)
        lines.append(text)
        lines.append("")
    return "\n".join(lines)


def export_ics(meeting: Meeting, action_items: List[ActionItem]) -> str:
    """
    Renders iCalendar (.ics) format with action items having due dates.
    """
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//MeetScribe//Meeting Action Items//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
    ]
    now_stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")

    items_with_dates = [it for it in action_items if it.due_date]
    if not items_with_dates:
        dt_start = (
            meeting.created_at.strftime("%Y%m%dT%H%M%SZ") if meeting.created_at else now_stamp
        )
        lines.extend(
            [
                "BEGIN:VEVENT",
                f"UID:meeting-{meeting.id}@meetscribe",
                f"DTSTAMP:{now_stamp}",
                f"DTSTART:{dt_start}",
                f"SUMMARY:{meeting.title}",
                f"DESCRIPTION:Notes for {meeting.title}",
                "STATUS:CONFIRMED",
                "END:VEVENT",
            ]
        )
    else:
        for item in items_with_dates:
            if not item.due_date:
                continue
            due_str = item.due_date.strftime("%Y%m%d")
            next_day = item.due_date.toordinal() + 1
            due_end_str = date.fromordinal(next_day).strftime("%Y%m%d")
            desc = (
                f"Action Item: {item.description}\\n"
                f"Meeting: {meeting.title}\\n"
                f"Owner: {item.owner or 'Unassigned'}\\n"
                f"Priority: {item.priority.capitalize()}\\n"
                f"Status: {item.status.capitalize()}"
            )
            lines.extend(
                [
                    "BEGIN:VEVENT",
                    f"UID:action-{item.id}@meetscribe",
                    f"DTSTAMP:{now_stamp}",
                    f"DTSTART;VALUE=DATE:{due_str}",
                    f"DTEND;VALUE=DATE:{due_end_str}",
                    f"SUMMARY:Action: {item.description}",
                    f"DESCRIPTION:{desc}",
                    "STATUS:CONFIRMED",
                    "END:VEVENT",
                ]
            )

    lines.append("END:VCALENDAR")
    return "\r\n".join(lines) + "\r\n"
